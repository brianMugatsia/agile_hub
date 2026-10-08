from datetime import timedelta
from io import BytesIO
from uuid import UUID
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from zipfile import ZipFile

from django import forms
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.db.models.functions import Lower
from django.http import HttpResponse, HttpResponseBadRequest
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone as django_timezone
from django.views import View
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.audit.services import record_event
from apps.hubs.permissions import hubs_for_user
from apps.inventory.models import InventoryBalance, InventoryTransaction
from apps.inventory.services import record_movement
from apps.products.forms import ProductForm
from apps.products.models import Product

from .models import StagedSpreadsheetImport


MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_EXPANDED_WORKBOOK_BYTES = 20 * 1024 * 1024
MAX_IMPORT_ROWS = 500
STAGE_LIFETIME = timedelta(minutes=15)


class ExcelUploadForm(forms.Form):
    file = forms.FileField()

    def clean_file(self):
        uploaded = self.cleaned_data["file"]
        if not uploaded.name.lower().endswith(".xlsx"):
            raise forms.ValidationError("Upload an .xlsx workbook.")
        if uploaded.size > MAX_IMPORT_BYTES:
            raise forms.ValidationError("The workbook must be 2 MB or smaller.")
        return uploaded


def read_workbook(uploaded, required_headers, allowed_headers):
    try:
        with ZipFile(uploaded) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_WORKBOOK_BYTES:
                raise ValidationError("The expanded workbook must be 20 MB or smaller.")
        uploaded.seek(0)
        workbook = load_workbook(uploaded, read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, ParseError) as error:
        raise ValidationError("The uploaded file is not a valid Excel workbook.") from error
    try:
        sheet = workbook.active
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = [
            str(cell.value or "").strip().lower().replace(" ", "_")
            for cell in header_cells
        ]
        non_empty_headers = [header for header in headers if header]
        if len(non_empty_headers) != len(set(non_empty_headers)):
            raise ValidationError("The workbook contains duplicate column headings.")
        unexpected = set(non_empty_headers) - set(allowed_headers)
        if unexpected:
            raise ValidationError(
                f"Unexpected columns: {', '.join(sorted(unexpected))}."
            )
        missing = set(required_headers) - set(non_empty_headers)
        if missing:
            raise ValidationError(
                f"Missing required columns: {', '.join(sorted(missing))}."
            )

        rows = []
        for line_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            if len(cells) > len(allowed_headers):
                raise ValidationError(f"Row {line_number} has too many columns.")
            if line_number > MAX_IMPORT_ROWS + 1:
                raise ValidationError(f"Upload no more than {MAX_IMPORT_ROWS} data rows.")
            if not any(cell.value not in (None, "") for cell in cells):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ValidationError(f"Upload no more than {MAX_IMPORT_ROWS} data rows.")
            if any(cell.data_type == "f" for cell in cells):
                raise ValidationError(f"Row {line_number} contains a formula; use values only.")
            values = {
                header: cells[index].value
                for index, header in enumerate(headers)
                if header and index < len(cells)
            }
            rows.append({"line": line_number, "values": values})
        if not rows:
            raise ValidationError("The workbook has no data rows.")
        return rows
    except (BadZipFile, InvalidFileException, OSError, ValueError, ParseError) as error:
        raise ValidationError("The uploaded file is not a valid Excel workbook.") from error
    finally:
        workbook.close()


class SpreadsheetImportView(LoginRequiredMixin, PermissionRequiredMixin, View):
    template_name = "imports/excel_upload.html"
    list_url = ""
    template_url = ""
    import_type = ""
    page_title = "Import Excel data"
    description = ""
    required_headers = ()
    template_headers = ()
    permission_required = ()

    def get_context_data(self, **kwargs):
        return {
            "page_title": self.page_title,
            "description": self.description,
            "template_url": self.template_url,
            "list_url": self.list_url,
            "preview_headers": self.template_headers,
            **kwargs,
        }

    def get(self, request):
        report_token = request.GET.get("error_report")
        if report_token:
            try:
                token = UUID(report_token)
            except (TypeError, ValueError, AttributeError):
                return HttpResponseBadRequest("This error report is invalid or expired.")
            staged = get_object_or_404(
                StagedSpreadsheetImport,
                token=token,
                import_type=self.import_type,
                created_by=request.user,
                report_only=True,
            )
            if django_timezone.now() - staged.created_at > STAGE_LIFETIME:
                staged.delete()
                return HttpResponseBadRequest("This error report expired. Upload the workbook again.")
            return self.error_report_response(staged)
        return render(
            request,
            self.template_name,
            self.get_context_data(form=ExcelUploadForm()),
        )

    def post(self, request):
        if request.POST.get("confirm_import") == "1":
            return self.confirm_import(request)

        StagedSpreadsheetImport.objects.filter(
            created_at__lt=django_timezone.now() - STAGE_LIFETIME,
        ).delete()
        form = ExcelUploadForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(
                request,
                self.template_name,
                self.get_context_data(form=form),
                status=400,
            )
        try:
            rows = read_workbook(
                form.cleaned_data["file"],
                self.required_headers,
                self.template_headers,
            )
            preview_rows, errors, import_rows = self.validate_rows(rows, request.user)
        except ValidationError as error:
            return render(
                request,
                self.template_name,
                self.get_context_data(
                    form=ExcelUploadForm(),
                    errors=error.messages,
                ),
                status=400,
            )
        context = self.get_context_data(
            form=ExcelUploadForm(),
            preview_rows=preview_rows,
            errors=errors,
        )
        if errors:
            row_errors = {}
            for error in errors:
                if error.startswith("Row "):
                    line_number, separator, message = error[4:].partition(": ")
                    if separator and line_number.isdigit():
                        row_errors.setdefault(int(line_number), []).append(message)
            rejected_rows = [
                {
                    "line": row["line"],
                    "values": {
                        header: str(row["values"].get(header, "") or "")
                        for header in self.template_headers
                    },
                    "errors": row_errors[row["line"]],
                }
                for row in rows
                if row["line"] in row_errors
            ]
            if rejected_rows:
                report = StagedSpreadsheetImport.objects.create(
                    import_type=self.import_type,
                    created_by=request.user,
                    report_only=True,
                    rows=rejected_rows,
                )
                context["error_report_url"] = (
                    f"{request.path}?error_report={report.token}"
                )
            return render(request, self.template_name, context, status=400)

        staged = StagedSpreadsheetImport.objects.create(
            import_type=self.import_type,
            created_by=request.user,
            rows=import_rows,
        )
        context.update(stage_token=staged.token, ready_to_import=True)
        return render(request, self.template_name, context)

    def error_report_response(self, staged):
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Rejected rows")
        sheet.append([*self.template_headers, "Errors"])
        for row in staged.rows:
            values = row["values"]
            cells = [values.get(header, "") for header in self.template_headers]
            cells.append("; ".join(row["errors"]))
            sheet.append([
                f"'{value}" if isinstance(value, str) and value and value[:1] in "=+-@" else value
                for value in cells
            ])
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="{self.import_type}-import-errors.xlsx"'
        )
        return response

    @transaction.atomic
    def confirm_import(self, request):
        try:
            token = UUID(request.POST.get("stage_token", ""))
        except (TypeError, ValueError, AttributeError):
            return HttpResponseBadRequest("This import preview is invalid or expired. Upload the workbook again.")
        staged = get_object_or_404(
            StagedSpreadsheetImport.objects.select_for_update(),
            token=token,
            import_type=self.import_type,
            created_by=request.user,
            report_only=False,
            consumed_at__isnull=True,
        )
        if django_timezone.now() - staged.created_at > STAGE_LIFETIME:
            staged.delete()
            return HttpResponseBadRequest("This import preview expired. Upload the workbook again.")

        row_count = len(staged.rows)
        try:
            with transaction.atomic():
                self.commit_rows(staged.rows, request.user, request=request)
        except ValidationError as error:
            messages.error(request, "Import was not applied: " + " ".join(error.messages))
            return redirect(request.path)
        except IntegrityError:
            staged.consumed_at = django_timezone.now()
            staged.rows = []
            staged.save(update_fields=["consumed_at", "rows"])
            messages.error(
                request,
                "Import was not applied because the data changed after preview. Upload and review the workbook again.",
            )
            return redirect(request.path)
        staged.consumed_at = django_timezone.now()
        staged.rows = []
        staged.save(update_fields=["consumed_at", "rows"])
        messages.success(request, f"Imported {row_count} rows.")
        return redirect(self.list_url)


class ExcelTemplateView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = ()
    headers = ()
    filename = "import-template.xlsx"

    def get(self, request):
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Import")
        sheet.append(self.headers)
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{self.filename}"'
        return response


def _preview_rows(rows, columns):
    return [
        [str(row["values"].get(key, "") or "") for key in columns]
        for row in rows[:20]
    ]


class ProductImportView(SpreadsheetImportView):
    import_type = StagedSpreadsheetImport.ImportType.PRODUCTS
    permission_required = "products.add_product"
    list_url = "products:list"
    template_url = "products:import_template"
    page_title = "Import products"
    description = "Create products from a validated Excel workbook."
    required_headers = ("sku", "name", "selling_price")
    template_headers = (
        "sku", "name", "selling_price", "category", "description", "unit",
        "cost_price", "reorder_level", "is_active",
    )

    def validate_rows(self, rows, user):
        errors = []
        accepted = []
        seen_skus = set()
        uploaded_skus = {
            str(row["values"].get("sku", "")).strip().casefold()
            for row in rows
            if row["values"].get("sku")
        }
        existing_skus = set(
            Product.objects.annotate(normalized_sku=Lower("sku"))
            .filter(normalized_sku__in=uploaded_skus)
            .values_list("normalized_sku", flat=True)
        )
        for row in rows:
            data = row["values"]
            for field, default in (
                ("cost_price", "0"),
                ("reorder_level", 0),
                ("unit", "each"),
                ("is_active", True),
            ):
                if data.get(field) in (None, ""):
                    data[field] = default
            sku = str(data.get("sku", "")).strip().upper()
            if not sku:
                errors.append(f"Row {row['line']}: SKU is required.")
                continue
            if sku.casefold() in seen_skus or sku.casefold() in existing_skus:
                errors.append(f"Row {row['line']}: SKU {sku} already exists or is duplicated in this workbook.")
                continue
            seen_skus.add(sku.casefold())
            data["sku"] = sku
            form = ProductForm(data)
            if not form.is_valid():
                detail = "; ".join(
                    f"{field}: {', '.join(field_errors)}"
                    for field, field_errors in form.errors.items()
                )
                errors.append(f"Row {row['line']}: {detail}")
                continue
            cleaned = form.cleaned_data
            cleaned["sku"] = cleaned["sku"].upper()
            accepted.append({
                key: str(value) if value is not None and key in ("cost_price", "selling_price") else value
                for key, value in cleaned.items()
            })
        preview = _preview_rows(rows, self.template_headers)
        return preview, errors, accepted

    def commit_rows(self, rows, user, *, request):
        for row in rows:
            form = ProductForm(row)
            if not form.is_valid():
                raise ValidationError("Product data changed after preview. Upload and review the workbook again.")
            product = form.save()
            record_event(
                action="products.product.imported",
                summary=f"Imported product {product}",
                actor=user,
                target=product,
                request=request,
            )


class ProductImportTemplateView(ExcelTemplateView):
    permission_required = "products.add_product"
    headers = ProductImportView.template_headers
    filename = "product-import-template.xlsx"


class InventoryImportView(SpreadsheetImportView):
    import_type = StagedSpreadsheetImport.ImportType.INVENTORY
    permission_required = "inventory.add_inventorytransaction"
    list_url = "inventory:list"
    template_url = "inventory:import_template"
    page_title = "Import stock movements"
    description = "Record stock receipts, issues, and adjustments through the inventory ledger."
    required_headers = ("hub_code", "product_sku", "kind", "quantity")
    template_headers = (
        "hub_code", "product_sku", "kind", "quantity", "direction",
        "unit_cost", "reference", "notes",
    )

    def validate_rows(self, rows, user):
        errors = []
        accepted = []
        hub_codes = {
            str(row["values"].get("hub_code", "")).strip().casefold()
            for row in rows
            if row["values"].get("hub_code")
        }
        product_skus = {
            str(row["values"].get("product_sku", "")).strip().casefold()
            for row in rows
            if row["values"].get("product_sku")
        }
        available_hubs = {
            hub.normalized_code: hub
            for hub in hubs_for_user(user, active_only=True)
            .annotate(normalized_code=Lower("code"))
            .filter(normalized_code__in=hub_codes)
        }
        products = {
            product.normalized_sku: product
            for product in Product.objects.filter(is_active=True)
            .annotate(normalized_sku=Lower("sku"))
            .filter(normalized_sku__in=product_skus)
        }
        candidates = []
        for row in rows:
            data = row["values"]
            hub = available_hubs.get(str(data.get("hub_code", "")).strip().casefold())
            product = products.get(str(data.get("product_sku", "")).strip().casefold())
            if hub is None:
                errors.append(f"Row {row['line']}: Hub is unknown, inactive, or outside your access.")
                continue
            if product is None:
                errors.append(f"Row {row['line']}: Active product SKU was not found.")
                continue
            kind = str(data.get("kind", "")).strip().upper()
            direction = str(data.get("direction", "") or "").strip().upper() or None
            allowed_kinds = {
                InventoryTransaction.Kind.RECEIPT,
                InventoryTransaction.Kind.ISSUE,
                InventoryTransaction.Kind.ADJUSTMENT,
            }
            if kind not in allowed_kinds:
                errors.append(f"Row {row['line']}: Choose a valid movement kind.")
                continue
            try:
                quantity_decimal = forms.DecimalField(min_value=1).clean(data.get("quantity"))
                quantity = int(quantity_decimal)
                if quantity_decimal != quantity:
                    raise ValueError
            except (TypeError, ValueError, ValidationError):
                errors.append(f"Row {row['line']}: Quantity must be a positive whole number.")
                continue
            if kind == InventoryTransaction.Kind.ADJUSTMENT and direction not in InventoryTransaction.Direction.values:
                errors.append(f"Row {row['line']}: Adjustments need direction IN or OUT.")
                continue
            expected = {
                InventoryTransaction.Kind.RECEIPT: InventoryTransaction.Direction.IN,
                InventoryTransaction.Kind.REFUND: InventoryTransaction.Direction.IN,
                InventoryTransaction.Kind.ISSUE: InventoryTransaction.Direction.OUT,
                InventoryTransaction.Kind.SALE: InventoryTransaction.Direction.OUT,
            }.get(kind)
            direction = direction or expected
            if expected and direction != expected:
                errors.append(f"Row {row['line']}: Movement kind {kind} requires direction {expected}.")
                continue
            if direction not in InventoryTransaction.Direction.values:
                errors.append(f"Row {row['line']}: Choose direction IN or OUT.")
                continue
            unit_cost = data.get("unit_cost")
            if unit_cost not in (None, ""):
                try:
                    unit_cost = str(forms.DecimalField(min_value=0, max_digits=12, decimal_places=2).clean(unit_cost))
                except ValidationError:
                    errors.append(f"Row {row['line']}: Unit cost must be a non-negative amount with at most two decimals.")
                    continue
            candidates.append({
                "line": row["line"],
                "pair": (hub.pk, product.pk),
                "row": {
                    "hub_id": str(hub.pk),
                    "product_id": product.pk,
                    "kind": kind,
                    "direction": direction,
                    "quantity": quantity,
                    "unit_cost": unit_cost,
                    "reference": str(data.get("reference", "") or "")[:80],
                    "notes": str(data.get("notes", "") or "")[:255],
                },
            })
        balances = {}
        pairs = list(dict.fromkeys(candidate["pair"] for candidate in candidates))
        for offset in range(0, len(pairs), 100):
            query = Q()
            for hub_id, product_id in pairs[offset:offset + 100]:
                query |= Q(hub_id=hub_id, product_id=product_id)
            if pairs[offset:offset + 100]:
                for hub_id, product_id, quantity in InventoryBalance.objects.filter(
                    query
                ).values_list("hub_id", "product_id", "quantity"):
                    balances[(hub_id, product_id)] = quantity
        projected = {}
        for candidate in candidates:
            pair = candidate["pair"]
            staged_row = candidate["row"]
            current_quantity = projected.get(pair, balances.get(pair, 0))
            quantity = staged_row["quantity"]
            if (
                staged_row["direction"] == InventoryTransaction.Direction.OUT
                and current_quantity < quantity
            ):
                errors.append(
                    f"Row {candidate['line']}: Only {current_quantity} are available at the selected hub."
                )
                continue
            projected[pair] = current_quantity + (
                quantity
                if staged_row["direction"] == InventoryTransaction.Direction.IN
                else -quantity
            )
            accepted.append(staged_row)
        return _preview_rows(rows, self.template_headers), errors, accepted

    def commit_rows(self, rows, user, *, request):
        hubs = {str(hub.pk): hub for hub in hubs_for_user(user, active_only=True)}
        products = {product.pk: product for product in Product.objects.filter(is_active=True)}
        for row in rows:
            hub = hubs.get(row["hub_id"])
            product = products.get(row["product_id"])
            if hub is None or product is None:
                raise ValidationError("A hub or product changed after preview. Upload and review the workbook again.")
            record_movement(
                hub=hub,
                product=product,
                kind=row["kind"],
                direction=row["direction"],
                quantity=row["quantity"],
                unit_cost=row["unit_cost"],
                reference=row["reference"],
                notes=row["notes"],
                actor=user,
            )


class InventoryImportTemplateView(ExcelTemplateView):
    permission_required = "inventory.add_inventorytransaction"
    headers = InventoryImportView.template_headers
    filename = "inventory-import-template.xlsx"

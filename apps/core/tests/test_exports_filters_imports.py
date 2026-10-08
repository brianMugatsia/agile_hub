from decimal import Decimal
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.accounts.roles import Role
from apps.core.models import SavedFilter, StagedSpreadsheetImport
from apps.hubs.models import Hub
from apps.inventory.models import InventoryBalance, InventoryTransaction
from apps.products.models import Product

pytestmark = pytest.mark.django_db


def _xlsx(headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return SimpleUploadedFile(
        "import.xlsx",
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def test_filtered_csv_and_excel_exports_use_scoped_list_query(client, make_user, roles):
    admin = make_user(role=Role.ADMIN)
    Product.objects.create(sku="EXP-1", name="Sunflower oil", selling_price=Decimal("125.00"))
    Product.objects.create(sku="EXP-2", name="Feed", selling_price=Decimal("50.00"))
    client.force_login(admin)

    csv_response = client.get(reverse("products:export", args=["csv"]), {"q": "sunflower"})
    csv_body = b"".join(csv_response.streaming_content)
    excel_response = client.get(reverse("products:export", args=["xlsx"]), {"q": "sunflower"})

    assert csv_response.status_code == 200
    assert b"Sunflower oil" in csv_body
    assert b"Feed" not in csv_body
    assert excel_response.status_code == 200
    assert excel_response["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    for route in (
        "sales:export",
        "inventory:export",
        "commissions:export",
        "payroll:salary_export",
    ):
        for file_format in ("csv", "xlsx"):
            response = client.get(reverse(route, args=[file_format]))
            assert response.status_code == 200
            if file_format == "xlsx":
                exported = load_workbook(BytesIO(response.content), read_only=True)
                assert next(exported.active.iter_rows(values_only=True), None) is not None
                exported.close()


def test_shared_search_matches_each_term_across_product_fields(client, make_user):
    admin = make_user(role=Role.ADMIN)
    matching = Product.objects.create(
        sku="SEARCH-1",
        name="Sunflower",
        category="Oilseed",
        selling_price=Decimal("125.00"),
    )
    Product.objects.create(
        sku="SEARCH-2",
        name="Sunflower feed",
        category="Livestock",
        selling_price=Decimal("50.00"),
    )
    client.force_login(admin)

    response = client.get(reverse("products:list"), {"q": "Sunflower Oilseed"})

    assert response.status_code == 200
    assert matching.name.encode() in response.content
    assert b"Sunflower feed" not in response.content


def test_export_denies_users_without_view_permission(client, make_user, roles):
    user = make_user(role=Role.WORKER)
    client.force_login(user)

    assert client.get(reverse("products:export", args=["csv"])).status_code == 403


def test_csv_export_neutralizes_spreadsheet_formulas(client, make_user, roles):
    admin = make_user(role=Role.ADMIN)
    Product.objects.create(
        sku="EXP-FORMULA",
        name="=HYPERLINK(\"https://example.invalid\")",
        selling_price=Decimal("10.00"),
    )
    client.force_login(admin)

    response = client.get(reverse("products:export", args=["csv"]))
    body = b"".join(response.streaming_content)

    assert b"'=HYPERLINK" in body


def test_excel_template_download_requires_permission_and_returns_workbook(
    client, make_user, roles
):
    admin = make_user(role=Role.ADMIN)
    client.force_login(admin)
    response = client.get(reverse("products:import_template"))

    assert response.status_code == 200
    assert response["Content-Disposition"].endswith('"product-import-template.xlsx"')
    assert response["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    worker = make_user(role=Role.WORKER)
    client.force_login(worker)
    assert client.get(reverse("products:import_template")).status_code == 403


def test_shared_filter_is_visible_on_matching_list_and_only_creator_can_delete(
    client, make_user, roles
):
    admin = make_user(role=Role.ADMIN)
    coworker = make_user(role=Role.ADMIN)
    client.force_login(admin)
    save_response = client.post(
        reverse("core:saved_filter_create"),
        {
            "model_key": "products.product",
            "name": "Active oils",
            "query_params": '{"q":["oil"]}',
            "return_to": reverse("products:list"),
        },
    )
    saved_filter = SavedFilter.objects.get()

    assert save_response.status_code == 302
    client.force_login(coworker)
    list_response = client.get(reverse("products:list"))
    delete_response = client.post(
        reverse("core:saved_filter_delete", args=[saved_filter.pk]),
        {"return_to": reverse("products:list")},
    )

    assert b"Active oils" in list_response.content
    assert b"?q=oil" in list_response.content
    assert delete_response.status_code == 404
    assert SavedFilter.objects.filter(pk=saved_filter.pk).exists()


def test_product_excel_upload_previews_then_creates_only_after_confirmation(
    client, make_user, roles
):
    admin = make_user(role=Role.ADMIN)
    client.force_login(admin)
    uploaded = _xlsx(
        ["sku", "name", "selling_price", "cost_price", "unit"],
        [["UP-1", "New sunflower oil", 80, 50, "bottle"]],
    )
    preview = client.post(
        reverse("products:import"),
        {"file": uploaded},
    )

    assert preview.status_code == 200
    assert b"New sunflower oil" in preview.content
    assert not Product.objects.filter(sku="UP-1").exists()
    staged = StagedSpreadsheetImport.objects.get(import_type="products")
    response = client.post(
        reverse("products:import"),
        {"confirm_import": "1", "stage_token": str(staged.token)},
    )

    assert response.status_code == 302
    product = Product.objects.get(sku="UP-1")
    assert product.name == "New sunflower oil"
    staged.refresh_from_db()
    assert staged.consumed_at is not None
    assert staged.rows == []


def test_invalid_product_workbook_does_not_stage_or_create_records(client, make_user, roles):
    admin = make_user(role=Role.ADMIN)
    client.force_login(admin)
    uploaded = _xlsx(
        ["sku", "name", "selling_price"],
        [["DUP-1", "One", 10], ["DUP-1", "Duplicate", 12]],
    )

    response = client.post(reverse("products:import"), {"file": uploaded})

    assert response.status_code == 400
    assert b"duplicated" in response.content
    assert not Product.objects.exists()
    report = StagedSpreadsheetImport.objects.get(report_only=True)
    assert not Product.objects.exists()

    download = client.get(
        reverse("products:import"),
        {"error_report": str(report.token)},
    )
    workbook = load_workbook(BytesIO(download.content), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()

    assert download.status_code == 200
    assert rows[0][-1] == "Errors"
    assert rows[1][0] == "DUP-1"
    assert "already exists or is duplicated" in rows[1][-1]
    assert len(rows) == 2


def test_import_error_report_is_owned_by_uploader(client, make_user, roles):
    uploader = make_user(role=Role.ADMIN)
    other_admin = make_user(role=Role.ADMIN)
    client.force_login(uploader)
    response = client.post(
        reverse("products:import"),
        {"file": _xlsx(
            ["sku", "name", "selling_price"],
            [["BAD-1", "", 10]],
        )},
    )
    report = StagedSpreadsheetImport.objects.get(report_only=True)

    client.force_login(other_admin)
    denied = client.get(
        reverse("products:import"),
        {"error_report": str(report.token)},
    )

    assert response.status_code == 400
    assert denied.status_code == 404


def test_inventory_import_respects_hub_access_and_records_ledger_movements(
    client, make_user, roles
):
    manager = make_user(role=Role.HUB_MANAGER)
    hub = Hub.objects.create(code="XLS-HUB", name="Spreadsheet Hub", manager=manager)
    out_of_scope_hub = Hub.objects.create(code="OUT-HUB", name="Other Hub")
    product = Product.objects.create(
        sku="XLS-PROD",
        name="Rice",
        selling_price=Decimal("3.00"),
        reorder_level=2,
    )
    client.force_login(manager)
    denied = client.post(
        reverse("inventory:import"),
        {"file": _xlsx(
            ["hub_code", "product_sku", "kind", "quantity"],
            [["OUT-HUB", product.sku, "RECEIPT", 5]],
        )},
    )

    assert denied.status_code == 400
    assert b"outside your access" in denied.content
    assert not InventoryTransaction.objects.exists()

    preview = client.post(
        reverse("inventory:import"),
        {"file": _xlsx(
            ["hub_code", "product_sku", "kind", "quantity"],
            [[hub.code, product.sku, "RECEIPT", 5]],
        )},
    )
    staged = StagedSpreadsheetImport.objects.get(import_type="inventory", report_only=False)
    assert preview.status_code == 200
    assert not InventoryBalance.objects.exists()
    response = client.post(
        reverse("inventory:import"),
        {"confirm_import": "1", "stage_token": str(staged.token)},
    )

    assert response.status_code == 302
    assert InventoryBalance.objects.get(hub=hub, product=product).quantity == 5
    assert InventoryTransaction.objects.filter(hub=hub, product=product).count() == 1
    assert out_of_scope_hub.pk != hub.pk


def test_inventory_import_rejects_outgoing_movements_beyond_projected_stock(
    client, make_user, roles
):
    manager = make_user(role=Role.HUB_MANAGER)
    hub = Hub.objects.create(code="STOCK-HUB", name="Stock Hub", manager=manager)
    product = Product.objects.create(
        sku="STOCK-PROD",
        name="Rice",
        selling_price=Decimal("3.00"),
    )
    InventoryBalance.objects.create(hub=hub, product=product, quantity=3)
    client.force_login(manager)

    response = client.post(
        reverse("inventory:import"),
        {"file": _xlsx(
            ["hub_code", "product_sku", "kind", "quantity"],
            [[hub.code, product.sku, "ISSUE", 2], [hub.code, product.sku, "ISSUE", 2]],
        )},
    )

    assert response.status_code == 400
    assert b"Only 1 are available" in response.content
    assert not StagedSpreadsheetImport.objects.filter(report_only=False).exists()
    assert not InventoryTransaction.objects.exists()
    assert InventoryBalance.objects.get(hub=hub, product=product).quantity == 3


def test_import_preview_token_is_single_use(client, make_user, roles):
    admin = make_user(role=Role.ADMIN)
    client.force_login(admin)
    client.post(
        reverse("products:import"),
        {"file": _xlsx(
            ["sku", "name", "selling_price"],
            [["ONCE-1", "Single use", 10]],
        )},
    )
    staged = StagedSpreadsheetImport.objects.get(import_type="products")
    confirm_url = reverse("products:import")
    payload = {"confirm_import": "1", "stage_token": str(staged.token)}

    assert client.post(confirm_url, payload).status_code == 302
    assert client.post(confirm_url, payload).status_code == 404
    assert Product.objects.filter(sku="ONCE-1").count() == 1

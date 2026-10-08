import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.http import HttpResponse, HttpResponseBadRequest, StreamingHttpResponse
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell import WriteOnlyCell


MAX_EXPORT_ROWS = 50000


class Echo:
    def write(self, value):
        return value


def _cell_value(row, field_path):
    parts = field_path.split("__")
    value = row
    for part in parts:
        if value is None:
            return ""
        value = getattr(value, part, "")
    if len(parts) == 1:
        display = getattr(row, f"get_{parts[0]}_display", None)
        if callable(display):
            value = display()
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float, bool)):
        return value
    return str(value)


def _spreadsheet_safe(value):
    if isinstance(value, (int, float, bool, Decimal)):
        return value
    value = str(value)
    return f"'{value}" if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


class ScopedModelExportView(LoginRequiredMixin, PermissionRequiredMixin, View):
    list_view_class = None
    permission_required = ()

    def get_permission_required(self):
        if self.permission_required:
            return super().get_permission_required()
        return (
            f"{self.list_view_class.model._meta.app_label}."
            f"view_{self.list_view_class.model._meta.model_name}",
        )

    def get(self, request, file_format):
        if file_format not in {"csv", "xlsx"}:
            return HttpResponseBadRequest("Choose csv or xlsx.")
        list_view = self.list_view_class()
        list_view.setup(request)
        queryset = list_view.get_queryset()
        if list_view.filter_error:
            return HttpResponseBadRequest(list_view.filter_error)
        columns = list_view.columns
        if not columns:
            return HttpResponseBadRequest("This list does not have export columns configured.")
        if queryset.count() > MAX_EXPORT_ROWS:
            return HttpResponseBadRequest(
                f"Export is limited to {MAX_EXPORT_ROWS} rows; narrow the filters and try again."
            )

        model = self.list_view_class.model
        filename = f"{model._meta.model_name}.{file_format}"
        if file_format == "csv":
            writer = csv.writer(Echo())

            def rows():
                yield writer.writerow([column["label"] for column in columns])
                for row in queryset.iterator(chunk_size=1000):
                    yield writer.writerow([
                        _spreadsheet_safe(_cell_value(row, column["field"]))
                        for column in columns
                    ])

            response = StreamingHttpResponse(rows(), content_type="text/csv; charset=utf-8")
        else:
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet(title=model._meta.verbose_name_plural[:31])
            sheet.freeze_panes = "A2"
            header = []
            for label in (column["label"] for column in columns):
                cell = WriteOnlyCell(sheet, value=label)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1D4ED8")
                header.append(cell)
            sheet.append(header)
            for row in queryset.iterator(chunk_size=1000):
                sheet.append([
                    _spreadsheet_safe(_cell_value(row, column["field"]))
                    for column in columns
                ])
            output = BytesIO()
            workbook.save(output)
            response = HttpResponse(
                output.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

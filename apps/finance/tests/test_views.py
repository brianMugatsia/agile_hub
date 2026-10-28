from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone

from apps.accounts.roles import Role
from apps.members.models import MemberProfile, Business
from apps.core.scoping import scope_queryset
from apps.finance.forms import CashFlowForm
from apps.finance.models import Asset, CashFlow, Expense, IncomeStatement, IncomeStatementItem
from apps.hubs.models import Hub, HubMembership
from apps.payroll.models import SalaryRecord, WorkerProfile
from apps.products.models import Product
from apps.sales.models import Sale, SaleItem

pytestmark = pytest.mark.django_db


@pytest.mark.parametrize(
    "route_name",
    ["finance:break_even", "finance:projections", "finance:ratios"],
)
def test_finance_calculator_pages_render_for_authorized_user(client, make_user, route_name):
    client.force_login(make_user(role=Role.SUPER_ADMIN))

    response = client.get(reverse(route_name))

    assert response.status_code == 200


def test_break_even_calculation_is_submitted_and_rendered(client, make_user):
    client.force_login(make_user(role=Role.SUPER_ADMIN))

    response = client.get(
        reverse("finance:break_even"),
        {
            "calculate": "1",
            "fixed_costs": "100.00",
            "selling_price": "20.00",
            "variable_cost_per_unit": "10.00",
        },
    )

    assert response.status_code == 200
    assert response.context["result"]["units"] == 10
    assert "Break-even revenue" in response.content.decode()


def test_projection_calculation_is_submitted_and_rendered(client, make_user):
    client.force_login(make_user(role=Role.SUPER_ADMIN))

    response = client.get(
        reverse("finance:projections"),
        {
            "calculate": "1",
            "starting_revenue": "100.00",
            "starting_expenses": "60.00",
            "revenue_growth_percent": "10",
            "expense_growth_percent": "5",
            "periods": "2",
        },
    )

    assert response.status_code == 200
    assert [period.revenue for period in response.context["periods"]] == [
        Decimal("110.00"),
        Decimal("121.00"),
    ]


def test_financial_overview_scopes_viewer_data_to_accessible_hubs(client, make_user):
    viewer = make_user(role=Role.VIEWER)
    agent = make_user(role=Role.SALES_AGENT)
    visible_hub = Hub.objects.create(code="FIN-VISIBLE", name="Visible finance hub")
    hidden_hub = Hub.objects.create(code="FIN-HIDDEN", name="Hidden finance hub")
    HubMembership.objects.create(hub=visible_hub, user=viewer)
    product = Product.objects.create(
        sku="FIN-COST", name="Finance cost item", selling_price=Decimal("20.00")
    )
    today = timezone.localdate()
    visible_sale = Sale.objects.create(
        hub=visible_hub,
        agent=agent,
        status=Sale.Status.COMPLETED,
        completed_at=timezone.now(),
        total=Decimal("25.00"),
    )
    hidden_sale = Sale.objects.create(
        hub=hidden_hub,
        agent=agent,
        status=Sale.Status.COMPLETED,
        completed_at=timezone.now(),
        total=Decimal("100.00"),
    )
    SaleItem.objects.create(
        sale=visible_sale, product=product, quantity=2, unit_price=Decimal("12.50"),
        unit_cost=Decimal("4.00"),
    )
    SaleItem.objects.create(
        sale=hidden_sale, product=product, quantity=10, unit_price=Decimal("10.00"),
        unit_cost=Decimal("9.00"),
    )
    visible_worker_user = make_user(role=Role.WORKER)
    hidden_worker_user = make_user(role=Role.WORKER)
    visible_worker = WorkerProfile.objects.create(
        user=visible_worker_user,
        hub=visible_hub,
        employee_code="FIN-VISIBLE-WORKER",
        job_title="Operator",
        monthly_salary=Decimal("300.00"),
        hired_on=today,
    )
    hidden_worker = WorkerProfile.objects.create(
        user=hidden_worker_user,
        hub=hidden_hub,
        employee_code="FIN-HIDDEN-WORKER",
        job_title="Operator",
        monthly_salary=Decimal("900.00"),
        hired_on=today,
    )
    for worker, amount in ((visible_worker, Decimal("300.00")), (hidden_worker, Decimal("900.00"))):
        SalaryRecord.objects.create(
            worker=worker,
            period_start=today,
            period_end=today,
            gross_amount=amount,
            status=SalaryRecord.Status.PENDING,
        )
    for hub, amount in ((visible_hub, Decimal("11.00")), (hidden_hub, Decimal("70.00"))):
        statement = IncomeStatement.objects.create(
            hub=hub,
            period_start=today,
            period_end=today,
        )
        IncomeStatementItem.objects.create(
            statement=statement,
            label="Operating cost",
            amount=amount,
            is_expense=True,
        )
    client.force_login(viewer)

    response = client.get(reverse("finance:income_statement"), {"start": today.isoformat()})

    assert response.status_code == 200
    assert response.context["period_start"] == today
    assert response.context["period_end"] == today
    assert response.context["revenue"] == Decimal("25.00")
    assert response.context["cost_of_goods"] == Decimal("8.00")
    assert response.context["payroll"] == Decimal("300.00")
    assert response.context["manual_expenses"] == Decimal("11.00")


@pytest.mark.parametrize(
    ("params", "message"),
    [
        ({"start": "2026-1-01"}, "start date"),
        ({"end": "2025-02-30"}, "end date"),
        ({"start": "2024-01-01", "end": "2025-01-01"}, "366 days"),
        ({"start": "2025-02-02", "end": "2025-02-01"}, "on or after"),
    ],
)
def test_financial_overview_rejects_invalid_periods(client, make_user, params, message):
    client.force_login(make_user(role=Role.SUPER_ADMIN))

    response = client.get(reverse("finance:income_statement"), params)

    assert response.status_code == 400
    assert message in response.content.decode().lower()


def test_hub_scoping_includes_business_linked_assets_for_assigned_hub(make_user):
    manager = make_user(role=Role.HUB_MANAGER)
    visible_hub = Hub.objects.create(code="ASSET-VISIBLE", name="Visible asset hub")
    hidden_hub = Hub.objects.create(code="ASSET-HIDDEN", name="Hidden asset hub")
    HubMembership.objects.create(hub=visible_hub, user=manager)
    member_user = make_user(role=Role.MEMBER)
    member = MemberProfile.objects.create(user=member_user)
    visible_business = Business.objects.create(
        member=member, hub=visible_hub, name="Visible business"
    )
    hidden_business = Business.objects.create(
        member=member, hub=hidden_hub, name="Hidden business"
    )
    visible_asset = Asset.objects.create(
        name="Business-only visible asset", cost=Decimal("100.00"), business=visible_business
    )
    Asset.objects.create(
        name="Hidden business-only asset", cost=Decimal("900.00"), business=hidden_business
    )
    Asset.objects.create(name="Hidden direct hub asset", cost=Decimal("300.00"), hub=hidden_hub)

    assert list(scope_queryset(Asset.objects.all(), manager)) == [visible_asset]


def test_cash_flow_form_rejects_business_from_different_hub(make_user):
    first_hub = Hub.objects.create(code="CASHFLOW-ONE", name="Cash flow hub one")
    second_hub = Hub.objects.create(code="CASHFLOW-TWO", name="Cash flow hub two")
    member_user = make_user(role=Role.MEMBER)
    member = MemberProfile.objects.create(user=member_user)
    business = Business.objects.create(member=member, hub=second_hub, name="Other hub business")

    form = CashFlowForm(data={
        "direction": CashFlow.Direction.INFLOW,
        "category": "Other income",
        "amount": "25.00",
        "transaction_date": "2026-10-09",
        "hub": str(first_hub.pk),
        "business": str(business.pk),
        "reference": "",
        "description": "",
    })

    assert not form.is_valid()
    assert "business" in form.errors


def test_cash_flow_create_limits_hubs_and_businesses_to_manager_scope(
    client, make_user, roles
):
    manager = make_user(role=Role.HUB_MANAGER)
    visible_hub = Hub.objects.create(code="CASHFLOW-VISIBLE", name="Visible cash flow hub")
    hidden_hub = Hub.objects.create(code="CASHFLOW-HIDDEN", name="Hidden cash flow hub")
    HubMembership.objects.create(hub=visible_hub, user=manager)
    member_user = make_user(role=Role.MEMBER)
    member = MemberProfile.objects.create(user=member_user)
    visible_business = Business.objects.create(
        member=member, hub=visible_hub, name="Visible cash flow business"
    )
    hidden_business = Business.objects.create(
        member=member, hub=hidden_hub, name="Hidden cash flow business"
    )
    client.force_login(manager)

    response = client.get(reverse("finance:cash_flow_create"))

    assert response.status_code == 200
    form = response.context["form"]
    assert list(form.fields["hub"].queryset) == [visible_hub]
    assert list(form.fields["business"].queryset) == [visible_business]
    assert hidden_business not in form.fields["business"].queryset


def test_expense_list_and_categories(client, make_user):
    admin = make_user(role=Role.SUPER_ADMIN)
    hub = Hub.objects.create(code="EXP-HUB", name="Expense hub")
    today = timezone.localdate()

    for category, _ in Expense.Category.choices:
        Expense.objects.create(
            hub=hub,
            category=category,
            title=f"Test {category}",
            amount=Decimal("100.00"),
            expense_date=today,
        )

    client.force_login(admin)
    response = client.get(reverse("finance:expenses"))
    assert response.status_code == 200
    assert len(response.context["object_list"]) == len(Expense.Category.choices)

    # Test filtering by category
    response_filtered = client.get(reverse("finance:expenses"), {"category": "RENT"})
    assert response_filtered.status_code == 200
    assert len(response_filtered.context["object_list"]) == 1
    assert response_filtered.context["object_list"][0].category == "RENT"


def test_expense_integration_with_income_statement(client, make_user):
    admin = make_user(role=Role.SUPER_ADMIN)
    hub = Hub.objects.create(code="EXP-IS-HUB", name="Expense IS hub")
    today = timezone.localdate()

    Expense.objects.create(
        hub=hub,
        category=Expense.Category.STOCK_COST_OF_GOODS,
        title="Stock purchase",
        amount=Decimal("150.00"),
        expense_date=today,
    )
    Expense.objects.create(
        hub=hub,
        category=Expense.Category.RENT,
        title="Office rent",
        amount=Decimal("500.00"),
        expense_date=today,
    )

    client.force_login(admin)
    response = client.get(reverse("finance:income_statement"), {"start": today.isoformat()})
    assert response.status_code == 200
    assert response.context["cost_of_goods"] == Decimal("150.00")
    assert response.context["manual_expenses"] == Decimal("500.00")
    assert response.context["operating_expenses"] == Decimal("500.00")


def test_expense_import_and_export(client, make_user):
    from io import BytesIO
    from django.core.files.uploadedfile import SimpleUploadedFile
    from openpyxl import Workbook, load_workbook

    admin = make_user(role=Role.SUPER_ADMIN)
    hub = Hub.objects.create(code="IMP-HUB", name="Import hub")
    client.force_login(admin)

    # 1. Test template download
    template_response = client.get(reverse("finance:expense_import_template"))
    assert template_response.status_code == 200

    # 2. Test Excel import upload and confirmation
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["hub_code", "category", "title", "amount", "expense_date", "reference", "notes"])
    sheet.append([hub.code, "RENT", "HQ Office Rent", "1200.00", "2026-10-01", "REF-001", "Monthly rent"])
    output = BytesIO()
    workbook.save(output)
    uploaded_file = SimpleUploadedFile(
        "expenses.xlsx",
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    # Post upload (preview)
    response = client.get(reverse("finance:import"))
    assert response.status_code == 200

    post_response = client.post(reverse("finance:import"), {"file": uploaded_file})
    assert post_response.status_code == 200
    assert "stage_token" in post_response.context

    stage_token = post_response.context["stage_token"]
    # Confirm import
    confirm_response = client.post(reverse("finance:import"), {"confirm_import": "1", "stage_token": stage_token})
    assert confirm_response.status_code == 302

    assert Expense.objects.filter(title="HQ Office Rent", amount=Decimal("1200.00")).exists()

    # 3. Test exports (CSV, Excel, PDF)
    for file_format in ("csv", "xlsx", "pdf"):
        export_response = client.get(reverse("finance:export", args=[file_format]))
        assert export_response.status_code == 200
        if file_format == "csv":
            assert b"HQ Office Rent" in b"".join(export_response.streaming_content)
        elif file_format == "xlsx":
            exported = load_workbook(BytesIO(export_response.content), read_only=True)
            assert next(exported.active.iter_rows(values_only=True), None) is not None
            exported.close()
        elif file_format == "pdf":
            assert b"HQ Office Rent" in export_response.content


def test_expense_deletion(client, make_user, roles):
    hub = Hub.objects.create(code="DEL-HUB", name="Delete Hub")
    admin = make_user(role=Role.ADMIN)
    expense = Expense.objects.create(
        hub=hub,
        category=Expense.Category.RENT,
        title="Test Rent",
        amount=Decimal("500.00"),
        expense_date=timezone.localdate(),
    )
    client.force_login(admin)
    response = client.post(reverse("finance:expense_delete", args=[expense.pk]))
    assert response.status_code == 302
    assert not Expense.objects.filter(pk=expense.pk).exists()

def test_expense_list_renders_add_modal_and_edit_buttons(client, make_user):
    hub = Hub.objects.create(code="LIST-HUB", name="List Hub")
    admin = make_user(role=Role.ADMIN)
    Expense.objects.create(
        hub=hub,
        category=Expense.Category.UTILITIES,
        title="Electricity Bill",
        amount=Decimal("350.00"),
        expense_date=timezone.localdate(),
    )
    client.force_login(admin)
    response = client.get(reverse("finance:expenses"))
    assert response.status_code == 200
    content = response.content.decode()
    assert "Add Single Expense" in content
    assert "Edit" in content
